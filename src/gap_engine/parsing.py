"""SOP document parsing.

Extracted from detector.py so app.py can count clauses before launching a scan
without importing chromadb and langchain.

Known limitation, deliberately not changed here: _parse_sop reads only
doc.paragraphs, so text in tables, headers and footers is never scanned. Many
pharmaceutical SOPs place acceptance criteria in tables. This is a scope limit
to disclose, not a silent behaviour change to make under demo pressure.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional


class ParseError(Exception):
    """The document could not be read at all."""


@dataclass
class SOPClause:
    section_id: str
    heading: str
    body: str


HEADING_STYLES = frozenset({
    "Heading 1", "Heading 2", "Heading 3",
    "heading 1", "heading 2", "heading 3",
})
SECTION_RE = re.compile(r'^(\d+(?:\.\d+)*\.?)\s+(.+)$')


def parse_docx_clauses(docx_path: Path) -> list[SOPClause]:
    import docx as _docx
    try:
        doc = _docx.Document(str(docx_path))
    except Exception as exc:
        raise ParseError(f"could not open {docx_path.name}: {type(exc).__name__}: {exc}") from exc

    clauses: list[SOPClause] = []
    current_section: Optional[str] = None
    current_heading: Optional[str] = None
    body_paras: list[str] = []

    def flush() -> None:
        if current_heading and body_paras:
            clauses.append(SOPClause(
                section_id=current_section or "§?",
                heading=current_heading,
                body=" ".join(body_paras),
            ))

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name if para.style else ""
        is_heading = style_name in HEADING_STYLES or bool(SECTION_RE.match(text))

        if is_heading:
            flush()
            m = SECTION_RE.match(text)
            if m:
                current_section = f"§{m.group(1).rstrip('.')}"
                current_heading = m.group(2).strip()
            else:
                current_section = "§?"
                current_heading = text
            body_paras = []
        else:
            body_paras.append(text)

    flush()
    return clauses


def parse_pdf_clauses(pdf_path: Path) -> list[SOPClause]:
    import fitz
    try:
        doc = fitz.open(str(pdf_path))
    except Exception as exc:
        raise ParseError(f"could not open {pdf_path.name}: {type(exc).__name__}: {exc}") from exc

    try:
        if getattr(doc, "needs_pass", False):
            raise ParseError(f"{pdf_path.name} is password-protected")
        lines: list[str] = []
        for page in doc:
            for line in page.get_text().splitlines():
                line = line.strip()
                if line:
                    lines.append(line)
    finally:
        doc.close()

    clauses: list[SOPClause] = []
    current_section: Optional[str] = None
    current_heading: Optional[str] = None
    body_lines: list[str] = []

    def flush() -> None:
        if current_heading and body_lines:
            clauses.append(SOPClause(
                section_id=current_section or "§?",
                heading=current_heading,
                body=" ".join(body_lines),
            ))

    for line in lines:
        m = SECTION_RE.match(line)
        if m:
            flush()
            current_section = f"§{m.group(1).rstrip('.')}"
            current_heading = m.group(2).strip()
            body_lines = []
        else:
            body_lines.append(line)

    flush()
    return clauses


def parse_clauses(path: Path) -> list[SOPClause]:
    """Dispatch on suffix. Raises ParseError on anything unreadable."""
    if path.suffix.lower() == ".pdf":
        return parse_pdf_clauses(path)
    return parse_docx_clauses(path)


def count_clauses(path: Path) -> int:
    """Cheap pre-check for the UI. Returns 0 if the document cannot be parsed."""
    try:
        return len(parse_clauses(path))
    except Exception:
        return 0
